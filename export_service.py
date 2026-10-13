"""
Export Service — CSV, Excel, PDF, JSON report generation
"""
import csv, json, io, datetime, os
from fpdf import FPDF

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except:
    HAS_OPENPYXL = False

# ─── CSV Export ────────────────────────────────────────────────────

def export_csv(rows, columns):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([c['label'] for c in columns])
    for row in rows:
        writer.writerow([str(row.get(c['key'], '') or '') for c in columns])
    return output.getvalue().encode('utf-8-sig')

# ─── JSON Export ───────────────────────────────────────────────────

def export_json(data):
    return json.dumps(data, ensure_ascii=False, indent=2, default=str).encode('utf-8')

# ─── Excel Export ──────────────────────────────────────────────────

def export_xlsx(rows, columns, title='Export'):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]

    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='1F2937', end_color='1F2937', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin', color='374151'),
        right=Side(style='thin', color='374151'),
        top=Side(style='thin', color='374151'),
        bottom=Side(style='thin', color='374151')
    )

    for ci, col in enumerate(columns, 1):
        cell = ws.cell(row=1, column=ci, value=col['label'])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for ri, row in enumerate(rows, 2):
        for ci, col in enumerate(columns, 1):
            val = row.get(col['key'], '')
            if isinstance(val, (list, dict)):
                val = json.dumps(val, ensure_ascii=False)
            cell = ws.cell(row=ri, column=ci, value=str(val or ''))
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')

    for ci, col in enumerate(columns, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = max(12, len(col['label']) + 4)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

# ─── PDF Export ────────────────────────────────────────────────────

class PDFReport(FPDF):
    def header(self):
        self.set_font('Helvetica', 'B', 10)
        self.set_text_color(156, 163, 175)
        self.cell(0, 8, 'Cyber Parc - Rapport d\'inventaire', align='L')
        self.ln(10)

    def footer(self):
        self.set_y(-15)
        self.set_font('Helvetica', 'I', 8)
        self.set_text_color(156, 163, 175)
        self.cell(0, 10, f'Page {self.page_no()}/{{nb}}', align='C')

def export_pdf(title, sections):
    pdf = PDFReport()
    pdf.alias_nb_pages()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=20)

    pdf.set_font('Helvetica', 'B', 18)
    pdf.set_text_color(31, 41, 55)
    pdf.cell(0, 15, title, align='L')
    pdf.ln(4)
    pdf.set_font('Helvetica', '', 9)
    pdf.set_text_color(107, 114, 128)
    pdf.cell(0, 8, f'Genere le {datetime.datetime.now().strftime("%d/%m/%Y %H:%M")}', align='L')
    pdf.ln(12)

    for section in sections:
        if pdf.get_y() > 240:
            pdf.add_page()

        pdf.set_font('Helvetica', 'B', 13)
        pdf.set_text_color(31, 41, 55)
        pdf.set_fill_color(243, 244, 246)
        pdf.cell(0, 10, section.get('title', ''), fill=True, ln=True)
        pdf.ln(4)

        if section.get('type') == 'kpi':
            pdf.set_font('Helvetica', '', 10)
            for kpi in section.get('data', []):
                pdf.set_text_color(55, 65, 81)
                pdf.cell(0, 7, f"  {kpi['label']}: {kpi['value']}", ln=True)
            pdf.ln(6)

        elif section.get('type') == 'table':
            cols = section.get('columns', [])
            data = section.get('data', [])
            col_widths = section.get('col_widths', None)

            if data and cols:
                if not col_widths:
                    col_widths = [max(190 // len(cols), 20) for _ in cols]
                available = sum(col_widths)
                if available > 190:
                    col_widths = [int(w * 190 / available) for w in col_widths]

                pdf.set_font('Helvetica', 'B', 8)
                pdf.set_fill_color(31, 41, 55)
                pdf.set_text_color(255, 255, 255)
                for ci, col in enumerate(cols):
                    pdf.cell(col_widths[ci], 7, col.get('label', ''), border=1, fill=True, align='C')
                pdf.ln()

                pdf.set_font('Helvetica', '', 7)
                pdf.set_text_color(55, 65, 81)
                for ri, row in enumerate(data):
                    if ri > 0 and ri % 30 == 0:
                        pdf.add_page()
                        pdf.set_font('Helvetica', 'B', 8)
                        pdf.set_fill_color(31, 41, 55)
                        pdf.set_text_color(255, 255, 255)
                        for ci, col in enumerate(cols):
                            pdf.cell(col_widths[ci], 7, col.get('label', ''), border=1, fill=True, align='C')
                        pdf.ln()
                        pdf.set_font('Helvetica', '', 7)
                        pdf.set_text_color(55, 65, 81)

                    if ri % 2 == 0:
                        pdf.set_fill_color(249, 250, 251)
                    else:
                        pdf.set_fill_color(255, 255, 255)

                    row_h = 6
                    for ci, col in enumerate(cols):
                        val = str(row.get(col.get('key', ''), '') or '')
                        if len(val) > 40:
                            val = val[:37] + '...'
                        pdf.cell(col_widths[ci], row_h, val, border=1, fill=True)
                    pdf.ln()
            pdf.ln(6)

        elif section.get('type') == 'text':
            pdf.set_font('Helvetica', '', 10)
            pdf.set_text_color(75, 85, 99)
            pdf.multi_cell(0, 6, section.get('content', ''))
            pdf.ln(4)

    return pdf.output(dest='S').encode('latin-1', errors='replace')

# ─── Column definitions ────────────────────────────────────────────

LOGICIELS_COLUMNS = [
    {'key': 'nom', 'label': 'Logiciel'},
    {'key': 'editeur', 'label': 'Editeur'},
    {'key': 'version', 'label': 'Version'},
    {'key': 'categorie', 'label': 'Categorie'},
    {'key': 'risque_niveau', 'label': 'Risque'},
    {'key': 'risque_score', 'label': 'Score'},
    {'key': 'eol', 'label': 'EOL'},
    {'key': 'install_source', 'label': 'Source'},
    {'key': 'nb_postes', 'label': 'Postes'},
]

MACHINES_COLUMNS = [
    {'key': 'nom_pc', 'label': 'Machine'},
    {'key': 'ip', 'label': 'IP'},
    {'key': 'utilisateur', 'label': 'Utilisateur'},
    {'key': 'os_nom', 'label': 'OS'},
    {'key': 'domaine', 'label': 'Domaine'},
    {'key': 'ram_gb', 'label': 'RAM (Go)'},
    {'key': 'score_securite', 'label': 'Score Securite'},
    {'key': 'score_conformite', 'label': 'Score Conformite'},
]

VULNS_COLUMNS = [
    {'key': 'cve_id', 'label': 'CVE'},
    {'key': 'logiciel_nom', 'label': 'Logiciel'},
    {'key': 'cvss_score', 'label': 'Score CVSS'},
    {'key': 'cvss_severity', 'label': 'Severite'},
    {'key': 'source', 'label': 'Source'},
    {'key': 'exploit_public', 'label': 'Exploit'},
    {'key': 'date_publiee', 'label': 'Date publication'},
]

LICENCES_COLUMNS = [
    {'key': 'nom', 'label': 'Logiciel'},
    {'key': 'editeur', 'label': 'Editeur'},
    {'key': 'categorie_licence', 'label': 'Type licence'},
    {'key': 'install_type', 'label': 'Installation'},
    {'key': 'nb_machines', 'label': 'Postes'},
    {'key': 'cout_annuel', 'label': 'Cout annuel'},
]
